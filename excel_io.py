"""
excel_io.py
Excel is used purely as a storage format for the backup snapshot and
the restore report - there is no interactive workbook/macro UI, and no
input tag list file (tags are discovered directly from OFS).
"""

from datetime import datetime
import openpyxl


def write_backup(path, rows):
    """
    Writes the backup snapshot workbook.
    rows: list of dicts with keys: tag, item_id, bind_ok, value,
    value_type, error_code, access_rights
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Backup"
    headers = ["TagName", "ItemID", "BindOK", "Value", "ValueType",
               "AccessRights", "ErrorCode", "BackupTimestamp"]
    ws.append(headers)

    timestamp = datetime.now().isoformat(timespec="seconds")
    for r in rows:
        ws.append([
            r["tag"],
            r["item_id"],
            r["bind_ok"],
            r["value"],
            r.get("value_type") or "",
            r.get("access_rights") or "",
            str(r.get("error_code", "")) if r.get("error_code") else "",
            timestamp,
        ])

    wb.save(path)


def read_backup(path):
    """
    Reads a backup workbook previously written by write_backup().
    Returns a list of dicts: {"tag": str, "value": <original type>,
    "value_type": str or None, "access_rights": str or None} for rows
    where BindOK was True at backup time (there is nothing useful to
    restore for a tag that never bound).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Backup"] if "Backup" in wb.sheetnames else wb.active

    rows = []
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: i for i, name in enumerate(header)}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or row[col["TagName"]] in (None, ""):
            continue
        if not row[col["BindOK"]]:
            continue  # nothing was actually backed up for this tag
        access_rights = row[col["AccessRights"]] if "AccessRights" in col else None
        value_type = row[col["ValueType"]] if "ValueType" in col else None
        rows.append({
            "tag": str(row[col["TagName"]]),
            "value": row[col["Value"]],
            "value_type": value_type or None,
            "access_rights": access_rights or None,
        })
    wb.close()
    return rows


def write_restore_report(path, rows):
    """
    Writes a report of what the restore step actually did, for audit
    purposes. rows: list of dicts with keys:
        tag, item_id, status, detail
    status is one of: "written", "skipped_readonly", "failed_bind",
    "failed_write"
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RestoreReport"
    ws.append(["TagName", "ItemID", "Status", "Detail", "Timestamp"])

    timestamp = datetime.now().isoformat(timespec="seconds")
    for r in rows:
        ws.append([r["tag"], r["item_id"], r["status"], r.get("detail", ""), timestamp])

    wb.save(path)
